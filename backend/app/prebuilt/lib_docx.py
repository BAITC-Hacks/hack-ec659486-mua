#!/usr/bin/env python3
"""lib_docx.py — детерминированное извлечение текста и структуры из ДИ (.docx).

Pass A пайплайна: из .docx достаём чистые абзацы (python-docx, с fallback на
zip+xml), режем на секции (Общие положения / Квалификационные требования /
Должностные обязанности / Права / Ответственность / Взаимодействие), внутри
обязанностей — на сферы («N. В сфере …:») и атомарные пункты («N)»/«N.»/буллеты).
Также собираем метаданные шапки: должность, код ДИ, утверждение, подчинение,
замещение.

Не галлюцинирует: только то, что есть в документе. LLM-обогащение — отдельный
Pass B (verb_canonical, work_nature, canonical_function_id).

Запуск как утилита:
  py scripts/lib_docx.py "<path.docx>"      # печать сегментированной структуры
"""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

# ── низкоуровневое извлечение абзацев ───────────────────────────────────────
# Каждый абзац → {"text": str, "list": bool}. Флаг list=True означает, что Word
# пометил абзац как элемент нумерованного/маркированного списка (numPr) — даже
# если в .text нет литерального «1)». Это критично: часть ДИ (напр. Пресс-служба)
# использует авто-нумерацию, и пункты обязанностей идут без печатных номеров.

# служебные «листы» в конце ДИ-шаблона (особенно в табличных макетах школ)
_FOOTER = {"лист согласования", "лист регистрации изменений", "лист ознакомления"}


def _norm0(s: str) -> str:
    return re.sub(r"\s+", " ", s.strip().lower()).strip(" .:")


def para_objs(path) -> list[dict]:
    path = str(path)
    try:
        import docx  # python-docx

        d = docx.Document(path)
        out = []
        for p in d.paragraphs:
            t = p.text.strip()
            if t and _norm0(t) not in _FOOTER:
                out.append({"text": t, "list": _is_list_item(p)})
        # Таблицы двух видов: (а) «контентная» — весь корпус ДИ внутри одной ячейки
        # многими абзацами (макет школ, ряда деканатов): отдаём КАЖДЫЙ абзац ячейки
        # отдельным объектом (с флагом списка), иначе обязанности схлопываются в один
        # блоб и fn=0; (б) «табличная» (Взаимодействие role|desc) — однострочные
        # ячейки склеиваем в строку через таб, как раньше. Объединённые ячейки
        # (merge) в row.cells повторяются — дедупим по id(cell._tc).
        for tbl in d.tables:
            for row in tbl.rows:
                seen_tc, singles, seen_txt = set(), [], set()
                for c in row.cells:
                    tc_id = id(c._tc)
                    if tc_id in seen_tc:
                        continue
                    seen_tc.add(tc_id)
                    cps = [pp for pp in c.paragraphs if pp.text.strip()]
                    if len(cps) >= 2:  # контентная ячейка
                        for pp in cps:
                            t = pp.text.strip()
                            if _norm0(t) not in _FOOTER:
                                out.append({"text": t, "list": _is_list_item(pp)})
                    elif len(cps) == 1:  # однострочная — в строку
                        t = cps[0].text.strip()
                        if t and t not in seen_txt and _norm0(t) not in _FOOTER:
                            seen_txt.add(t)
                            singles.append(t)
                if singles:
                    out.append({"text": "\t".join(singles), "list": False})
        return out
    except Exception:
        return _paragraphs_zip(path)


def _is_list_item(p) -> bool:
    try:
        pPr = p._p.pPr
        if pPr is not None and pPr.numPr is not None:
            return True
    except Exception:
        pass
    return False


def paragraphs(path) -> list[str]:
    """Плоский список текстов абзацев (для harvest_meta)."""
    return [o["text"] for o in para_objs(path)]


def _paragraphs_zip(path) -> list[dict]:
    z = zipfile.ZipFile(path)
    xml = z.read("word/document.xml").decode("utf-8", "ignore")
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<w:tab[^>]*/>", "\t", xml)
    xml = re.sub(r"<[^>]+>", "", xml)
    return [{"text": ln.strip(), "list": False} for ln in xml.split("\n") if ln.strip()]


# ── распознавание заголовков секций ─────────────────────────────────────────
SECTIONS = {
    "general": ["общие положения"],
    "qualification": [
        "квалификационные требования",
        "квалификационным требованиям",
        "требования к квалификации",
    ],
    "duties": [
        "должностные обязанности",
        "функциональные обязанности",
        "должностные обязанности и функции",
    ],
    "rights": ["права"],
    "responsibility": ["ответственность"],
    "interaction": [
        "взаимодействие",
        "взаимоотношения",
        "служебные взаимоотношения",
        "связи по должности",
    ],
}
# заголовок секции = (опц. номер) + название, коротко (без двоеточия-перечня)
_SECTION_NUM = re.compile(r"^\s*(\d+)\.?\s*")


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.strip().lower()).strip(" .:")


def match_section(line: str):
    """Заголовок секции = короткая самостоятельная строка, почти точно совпадающая
    с названием. Иначе пункт-список «Взаимодействие со структурными…» ложно
    обрезал бы секцию обязанностей."""
    body = _SECTION_NUM.sub("", line).strip()
    n = _norm(body)
    if len(n) > 45:
        return None
    for key, variants in SECTIONS.items():
        for v in variants:
            if n == v or (n.startswith(v) and len(n) <= len(v) + 12):
                return key
    return None


# сфера внутри обязанностей. Номер может отсутствовать (директор ДАиУК:
# «К должностным обязанностям директора … относятся:» без номера и в виде списка).
_SPHERE_DUTY = re.compile(
    r"к (?:должностным|функциональным) обязанностям .*относятся$", re.IGNORECASE
)
_SPHERE_AREA = re.compile(
    r"(?:в сфере|в области|в части|в рамках|в системе|по вопросам|по направлению) .+", re.IGNORECASE
)
_SPHERE_KNOW = re.compile(r"должен (?:знать|уметь)\b", re.IGNORECASE)


def match_sphere(line: str):
    """→ (kind, label) где kind ∈ {duty, area, know}; иначе None.
    Анкеры «в сфере/…» требуют либо номер, либо двоеточие в конце — чтобы не
    ловить пункты-обязанности, начинающиеся со слов «в области …»."""
    m = re.match(r"^\s*(\d+)\.\s*(.+)$", line)
    num = m.group(1) if m else ""
    body = (m.group(2) if m else line).strip()
    bl = _norm(body)
    ends_colon = line.rstrip().endswith(":")
    label = body.rstrip(": ").strip()
    if _SPHERE_DUTY.search(bl):
        return ("duty", f"{num + '. ' if num else ''}{label}")
    if _SPHERE_KNOW.match(bl) and (num or ends_colon):
        return ("know", f"{num + '. ' if num else ''}{label}")
    if _SPHERE_AREA.match(bl) and (num or ends_colon):
        return ("area", f"{num + '. ' if num else ''}{label}")
    return None


# атомарный пункт: «1)», «1.», «1.1)», «-», «•», «–»
ITEM_RE = re.compile(r"^\s*(\d+[.)]|\d+\.\d+[.)]?|[-–•·▪])\s+(.*)$")
ENUM_ONLY = re.compile(r"^\s*(\d+[.)]|[-–•·▪])\s*$")


# ── метаданные шапки ────────────────────────────────────────────────────────
DOC_CODE_RE = re.compile(r"(ДИ[-\s]?[А-ЯA-Z]{0,4}[-\s]?\d+\.\d+/\d+[-–]\d+)")
APPROVAL_RE = re.compile(r"(Утвержден[ао].{0,80}?(?:Приказ|приказом).{0,80})", re.IGNORECASE)
REPORTS_RE = re.compile(r"подчин\w+\s+(?:непосредственно\s+)?([^.,;\n]{3,80})", re.IGNORECASE)
SUBST_RE = re.compile(
    r"обязанности\s+(?:возлагаются|исполня\w+)\s+(?:на|)\s*([^.,;\n]{3,90})", re.IGNORECASE
)


def harvest_meta(paras: list[str]) -> dict:
    text = "\n".join(paras)
    meta = {
        "doc_code": "",
        "approval": "",
        "reports_to": "",
        "substituted_by": "",
        "title_guess": "",
    }
    m = DOC_CODE_RE.search(text)
    if m:
        meta["doc_code"] = m.group(1).strip()
    m = APPROVAL_RE.search(text)
    if m:
        meta["approval"] = re.sub(r"\s+", " ", m.group(1)).strip()
    m = REPORTS_RE.search(text)
    if m:
        meta["reports_to"] = m.group(1).strip()
    m = SUBST_RE.search(text)
    if m:
        meta["substituted_by"] = m.group(1).strip()
    # заголовок-должность: первая строка с «специалист/директор/…» рядом с кодом ДИ,
    # либо строка после слова ДОЛЖНОСТНАЯ ИНСТРУКЦИЯ
    RANKS = (
        "директор",
        "руководител",
        "начальник",
        "специалист",
        "методист",
        "менеджер",
        "инженер",
        "секретар",
        "оператор",
        "фотограф",
        "дизайнер",
        "заведующ",
        "лаборант",
        "техник",
        "экономист",
        "бухгалтер",
        "юрист",
    )
    for p in paras[:40]:
        pl = p.lower()
        if any(r in pl for r in RANKS) and 4 <= len(p) <= 90 and "должностная инструкц" not in pl:
            meta["title_guess"] = p.strip()
            break
    return meta


# ── основной сегментатор ────────────────────────────────────────────────────
def segment(path) -> dict:
    objs = para_objs(path)
    objs = _drop_toc(objs)
    meta = harvest_meta([o["text"] for o in objs])

    sections: dict[str, list[dict]] = {}
    cur = None
    for o in objs:
        sec = match_section(o["text"])
        if sec:
            cur = sec
            sections.setdefault(cur, [])
            continue
        if cur:
            sections.setdefault(cur, []).append(o)

    duties, knowledge = _parse_duties(sections.get("duties", []))
    result = {
        "source": str(path),
        "meta": meta,
        "n_paragraphs": len(objs),
        "sections_found": sorted(sections.keys()),
        "general": [o["text"] for o in sections.get("general", [])],
        "qualification": [o["text"] for o in sections.get("qualification", [])],
        "duties": duties,
        "knowledge": knowledge,
        "rights": _items_only(sections.get("rights", [])),
        "responsibility": _items_only(sections.get("responsibility", [])),
        "interaction": [o["text"] for o in sections.get("interaction", [])],
    }
    return result


def _drop_toc(objs: list[dict]) -> list[dict]:
    out, skip = [], False
    for o in objs:
        if _norm(o["text"]) == "содержание":
            skip = True
            continue
        if skip:
            if match_section(o["text"]) == "general" or _norm(o["text"]).startswith(
                "общие положения"
            ):
                skip = False
                out.append(o)
            continue
        out.append(o)
    return out


def _parse_duties(objs: list[dict]) -> tuple[list[dict], list[str]]:
    """Режем обязанности на сферы и атомарные пункты.

    Пункт = литеральный «N)»/буллет ИЛИ абзац-список (numPr). Абзац без
    нумерации и без list-флага считается продолжением (перенос строки).
    Сферы «Должен знать/уметь» (компетенции, не функции) выносятся отдельно.
    Возвращает (duty_spheres, knowledge_items).
    """
    spheres: list[dict] = []
    knowledge: list[str] = []
    cur = {"sphere": "Должностные обязанности (общий перечень)", "items": []}
    cur_kind = "duty"
    seq = 0

    def flush():
        nonlocal cur
        if cur["items"]:
            if cur_kind == "know":
                knowledge.extend(it["text"] for it in cur["items"])
            else:
                spheres.append(cur)

    def add_item(num, text):
        nonlocal seq
        seq += 1
        cur["items"].append({"num": str(num), "text": text})

    for o in objs:
        ln, is_list = o["text"], o["list"]
        sp = match_sphere(ln)
        if sp:
            flush()
            kind, label = sp
            cur = {"sphere": label[0].upper() + label[1:] if label else label, "items": []}
            cur_kind = kind
            seq = 0
            continue
        mi = ITEM_RE.match(ln)
        if mi:
            add_item(mi.group(1).rstrip(".)"), mi.group(2).strip())
        elif ENUM_ONLY.match(ln):
            continue
        elif is_list:  # авто-нумерованный пункт без печатного №
            add_item(seq + 1, ln.strip())
        elif cur["items"]:  # продолжение предыдущего пункта (перенос)
            cur["items"][-1]["text"] = (cur["items"][-1]["text"] + " " + ln.strip()).strip()
        # иначе: вводный абзац до первого пункта — игнор
    flush()
    return spheres, knowledge


def _items_only(objs: list[dict]) -> list[str]:
    out, pending = [], None
    for o in objs:
        ln, is_list = o["text"], o["list"]
        mi = ITEM_RE.match(ln)
        if mi:
            if pending:
                out.append(pending)
            pending = mi.group(2).strip()
        elif is_list:
            if pending:
                out.append(pending)
            pending = ln.strip()
        elif pending:
            pending = (pending + " " + ln.strip()).strip()
    if pending:
        out.append(pending)
    return out


# ── CLI ─────────────────────────────────────────────────────────────────────
def _main(argv):
    if not argv:
        print("usage: py scripts/lib_docx.py <path.docx>", file=sys.stderr)
        return 2
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8")
        except Exception:
            pass
    r = segment(argv[0])
    m = r["meta"]
    print(f"FILE: {Path(argv[0]).name}")
    print(f"  paragraphs: {r['n_paragraphs']}  sections: {r['sections_found']}")
    print(f"  title_guess: {m['title_guess']}")
    print(f"  doc_code: {m['doc_code']}")
    print(f"  approval: {m['approval'][:90]}")
    print(f"  reports_to: {m['reports_to']}")
    print(f"  substituted_by: {m['substituted_by']}")
    print(f"  qualification items: {len(r['qualification'])}")
    nfun = sum(len(s["items"]) for s in r["duties"])
    print(f"  DUTIES: {len(r['duties'])} spheres, {nfun} atomic functions")
    for sph in r["duties"]:
        print(f"    ── [{sph['sphere']}]  ({len(sph['items'])})")
        for it in sph["items"][:3]:
            print(f"        {it['num']}) {it['text'][:96]}")
        if len(sph["items"]) > 3:
            print(f"        … +{len(sph['items']) - 3} ещё")
    print(f"  knowledge(должен знать): {len(r['knowledge'])}")
    print(
        f"  rights: {len(r['rights'])}  responsibility: {len(r['responsibility'])}"
        f"  interaction-lines: {len(r['interaction'])}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(_main(sys.argv[1:]))
